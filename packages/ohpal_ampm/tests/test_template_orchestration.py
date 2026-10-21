"""Tests for ohpal.ampm.template.make_template (the one-call wrapper)."""

from __future__ import annotations

import json
from pathlib import Path

import polars as pl
import pytest

from ohpal.ampm.template import make_template
from ohpal.samples.io import read_table

REAL_CSV = Path("/mnt/user-data/uploads/JR299_sterling_parts.csv")
REAL_DHXML = Path("/mnt/user-data/uploads/BuildStarted_2025_12_2_13_37_55.dhxml")


def _dhxml(tmp_path: Path) -> Path:
    doc = {"version1": {"build": {"mttFileName": "JobZ_OHP_Bar.mtt", "name": "", "parts": [
        {"name": "1", "boundingBox": "0,0,0,10,10,13"},
        {"name": "2", "boundingBox": "11,0,0,21,10,13"},
    ]}}}
    p = tmp_path / "BuildStarted.dhxml"
    p.write_text(json.dumps(doc), encoding="utf-8")
    return p


def test_without_quant_builds_plain_template(tmp_path):
    out = make_template(_dhxml(tmp_path), tmp_path / "t.xlsx")
    back = read_table(out, anchor="part_id")
    assert back["part_id"].to_list() == ["1", "2"]
    assert back["build_id"].to_list() == ["JobZ_OHP_Bar", "JobZ_OHP_Bar"]  # mtt-derived
    # no parameter columns when quant omitted
    assert "speed" not in back.columns and "Hatch Speed" not in back.columns


def test_template_kwargs_forwarded(tmp_path):
    out = make_template(_dhxml(tmp_path), tmp_path / "t.xlsx",
                        sheet_name="MySheet", include_bounds=False)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["MySheet"]
    assert "x_min" not in [c.value for c in wb.active[1]]  # bounds suppressed


@pytest.mark.skipif(not (REAL_CSV.exists() and REAL_DHXML.exists()),
                    reason="real QuantAM CSV + dhxml not present")
def test_with_quant_appends_parameters(tmp_path):
    out = make_template(
        REAL_DHXML, tmp_path / "jr299.xlsx",
        quant=REAL_CSV,
        rename={"Hatch Speed": "speed", "Hatches Power": "power"},
    )
    back = read_table(out, anchor="part_id")
    assert back.height == 20
    assert "speed" in back.columns and "power" in back.columns
    # part 1 is the low-speed corner of the parameter grid
    row = back.filter(pl.col("part_id") == "1")
    assert row["speed"].to_list() == [500.0]
    assert row["power"].to_list() == [150.0]
