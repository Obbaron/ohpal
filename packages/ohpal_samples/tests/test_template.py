"""Tests for the template generator and workbook writer."""

from __future__ import annotations

import json
from pathlib import Path

import polars as pl
import pytest
from openpyxl import load_workbook

from ohpal.core.build import BuildParts
from ohpal.samples.io import read_table
from ohpal.samples.template import (
    build_template_frame,
    create_template,
    create_template_from_dhxml,
)
from ohpal.samples.writer import WriteError, write_table

REAL_DHXML = Path("/mnt/user-data/uploads/BuildStarted_2025_12_2_13_37_55.dhxml")


def _build_with_mtt(tmp_path: Path, mtt="JobZ_OHP_Bar.mtt") -> BuildParts:
    doc = {"version1": {"build": {"mttFileName": mtt, "name": "", "parts": [
        {"name": "1", "boundingBox": "0,0,0,10,10,13"},
    ]}}}
    p = tmp_path / "BuildStarted.dhxml"
    p.write_text(json.dumps(doc), encoding="utf-8")
    return BuildParts.from_dhxml(p)


def _build(tmp_path: Path) -> BuildParts:
    doc = {"version1": {"build": {"parts": [
        {"name": "1", "boundingBox": "0,0,0,10,10,13"},
        {"name": "2", "boundingBox": "11,0,0,21,10,13"},
        {"name": "1", "boundingBox": "0,11,0,10,21,13"},  # duplicate name -> suffixed
    ]}}}
    p = tmp_path / "BuildStarted.dhxml"
    p.write_text(json.dumps(doc), encoding="utf-8")
    return BuildParts.from_dhxml(p)


# ---- frame assembly -------------------------------------------------------

def test_frame_has_keys_alias_and_coords(tmp_path):
    f = build_template_frame(_build(tmp_path), build_id="b1")
    assert f.columns == ["build_id", "part_id", "alias", "x_min", "x_max", "y_min", "y_max", "z_min", "z_max"]
    assert f["build_id"].to_list() == ["b1", "b1", "b1"]
    assert f["part_id"].to_list() == ["1#1", "2", "1#2"]   # dedup-suffix applied
    assert f["alias"].null_count() == 3                   # empty for the human
    assert f["x_min"].to_list()[1] == 11.0 and f["x_max"].to_list()[1] == 21.0  # part 2 bounds


def test_coords_and_cube_id_optional(tmp_path):
    f = build_template_frame(_build(tmp_path), build_id="b1",
                             include_bounds=False, include_alias=False)
    assert f.columns == ["build_id", "part_id"]


def test_extra_frame_joins_on_part_id(tmp_path):
    extra = pl.DataFrame({"part_id": ["2"], "ampm__speed": [800.0]})
    f = build_template_frame(_build(tmp_path), build_id="b1", extra=extra)
    assert "ampm__speed" in f.columns
    row = f.filter(pl.col("part_id") == "2")
    assert row["ampm__speed"].to_list() == [800.0]
    # a part absent from extra gets null, not dropped
    assert f.height == 3


def test_extra_without_part_id_raises(tmp_path):
    with pytest.raises(ValueError, match="part_id"):
        build_template_frame(_build(tmp_path), build_id="b1",
                             extra=pl.DataFrame({"speed": [1.0]}))


# ---- writing --------------------------------------------------------------

def test_create_xlsx_marks_generated_and_input_columns(tmp_path):
    out = create_template(_build(tmp_path), tmp_path / "t.xlsx", build_id="b1")
    wb = load_workbook(out)
    ws = wb.active
    headers = {c.value: c for c in ws[1]}
    # generated columns shaded grey, alias shaded yellow
    assert headers["part_id"].fill.fgColor.rgb.endswith("D9D9D9")
    assert headers["alias"].fill.fgColor.rgb.endswith("FFF2CC")
    # legend present as header comments
    assert headers["part_id"].comment is not None
    assert ws.freeze_panes == "A2"


def test_sheet_named_after_build_id(tmp_path):
    out = create_template(_build(tmp_path), tmp_path / "t.xlsx", build_id="Ti-Sterling")
    assert load_workbook(out).sheetnames == ["Ti-Sterling"]


def test_csv_output_is_plain(tmp_path):
    out = create_template(_build(tmp_path), tmp_path / "t.csv", build_id="b1")
    txt = out.read_text()
    assert "part_id" in txt.splitlines()[0]


def test_unsupported_output_type_raises(tmp_path):
    with pytest.raises(WriteError):
        write_table(pl.DataFrame({"a": [1]}), tmp_path / "x.parquet")


# ---- round-trip: generate -> read back with our own reader ----------------

def test_generated_template_reads_back(tmp_path):
    build = _build(tmp_path)
    out = create_template(build, tmp_path / "t.xlsx", build_id="b1")
    # part_id is the populated, authoritative anchor for a generated template
    back = read_table(out, anchor="part_id")
    assert back["part_id"].to_list() == build.part_ids
    assert set(back.columns) >= {"build_id", "part_id", "alias", "x_min", "x_max", "y_min", "y_max", "z_min", "z_max"}
    assert back.schema["part_id"] == pl.String


@pytest.mark.skipif(not REAL_DHXML.exists(), reason="build file not present")
def test_real_dhxml_round_trip(tmp_path):
    out = create_template_from_dhxml(REAL_DHXML, tmp_path / "real.xlsx")
    back = read_table(out, anchor="part_id")
    build = BuildParts.from_dhxml(REAL_DHXML)
    assert back["build_id"][0] == "JR299_OHP_Ti-Sterling"  # derived from mtt
    assert back.height == len(build) == 20
    assert sorted(back["part_id"].to_list(), key=int) == sorted(build.part_ids, key=int)
    assert back["alias"].null_count() == 20      # nothing filled in yet
    # coordinates survived as numbers
    assert back.schema["x_min"] == pl.Float64


# ---- build_id defaulting from the mtt filename ----------------------------

def test_build_id_defaults_to_mtt_stem(tmp_path):
    build = _build_with_mtt(tmp_path, "JobZ_OHP_Bar.mtt")
    out = create_template(build, tmp_path / "t.xlsx")  # no build_id supplied
    back = read_table(out, anchor="part_id")
    assert back["build_id"].to_list() == ["JobZ_OHP_Bar"]
    # sheet also named from the derived id
    from openpyxl import load_workbook
    assert load_workbook(out).sheetnames == ["JobZ_OHP_Bar"]


def test_explicit_build_id_overrides_mtt(tmp_path):
    build = _build_with_mtt(tmp_path, "JobZ_OHP_Bar.mtt")
    out = create_template(build, tmp_path / "t.xlsx", build_id="override")
    assert read_table(out, anchor="part_id")["build_id"].to_list() == ["override"]


def test_missing_build_id_raises(tmp_path):
    # _build has no mtt filename and no name -> nothing to derive from
    with pytest.raises(ValueError, match="build_id"):
        create_template(_build(tmp_path), tmp_path / "t.xlsx")


def test_bounds_grouped_into_outline(tmp_path):
    import zipfile, re
    build = _build(tmp_path)
    out = create_template(build, tmp_path / "g.xlsx", build_id="b1")
    with zipfile.ZipFile(out) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    # the six bound columns (D..I) carry an outline level
    cols = re.findall(r'<col[^>]*outlineLevel="1"[^>]*>', xml)
    assert cols, "no grouped columns found"
    assert 'min="4"' in cols[0] and 'max="9"' in cols[0]  # x_min..z_max


def test_non_contiguous_group_rejected(tmp_path):
    import pytest
    from ohpal.samples.writer import write_xlsx, WriteError
    df = pl.DataFrame({"part_id": ["1"], "a": [1.0], "b": [2.0], "c": [3.0]})
    with pytest.raises(WriteError, match="contiguous"):
        write_xlsx(df, tmp_path / "x.xlsx", column_groups=[["a", "c"]])  # b is between
